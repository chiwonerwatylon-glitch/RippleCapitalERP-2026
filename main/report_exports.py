"""Excel, CSV and DOCX exports for report payloads built by REPORT_BUILDERS."""
import csv
import io
import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Pt, RGBColor

COMPANY_NAME = "Ripple Capital Insurance"
BRAND_HEX = "1B702D"

REPORT_FORMATS = {
    "pdf": ("application/pdf", ".pdf"),
    "excel": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"),
    "csv": ("text/csv", ".csv"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
}


def _plain(value):
    return "" if value is None else value


def _text(value):
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:,.2f}"
    return str(value)


def build_report_csv(report):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(report["headers"])
    for row in report["rows"]:
        writer.writerow([_plain(v) for v in row])
    return BytesIO(buf.getvalue().encode("utf-8-sig"))


def build_report_xlsx(report):
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", report["title"])[:31] or "Report"

    ws.append([COMPANY_NAME])
    ws["A1"].font = Font(bold=True, size=14, color=BRAND_HEX)
    ws.append([report["title"]])
    ws["A2"].font = Font(bold=True, size=12)
    meta = f"Generated {datetime.now().strftime('%d %b %Y, %H:%M')}"
    if report.get("subtitle"):
        meta += f"  ·  {report['subtitle']}"
    ws.append([meta])
    for label, value in report.get("summary") or []:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(list(report["headers"]))
    fill = PatternFill("solid", fgColor=BRAND_HEX)
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    numeric = set(report.get("numeric_cols", ()))
    for row in report["rows"]:
        ws.append([_plain(v) for v in row])
        for idx in numeric:
            if idx < len(row):
                ws.cell(row=ws.max_row, column=idx + 1).number_format = "#,##0.00"

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for idx, header in enumerate(report["headers"], start=1):
        longest = max(
            [len(str(header))] + [len(_text(r[idx - 1])) for r in report["rows"] if idx - 1 < len(r)]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 45)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_report_docx(report):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    if section.page_width < section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width

    brand = doc.add_paragraph()
    run = brand.add_run(COMPANY_NAME)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(BRAND_HEX)

    doc.add_heading(report["title"], level=1)
    meta = f"Generated {datetime.now().strftime('%d %b %Y, %H:%M')}"
    if report.get("subtitle"):
        meta += f"  ·  {report['subtitle']}"
    doc.add_paragraph(meta)

    if report.get("summary"):
        summary_table = doc.add_table(rows=0, cols=2)
        summary_table.style = "Table Grid"
        for label, value in report["summary"]:
            cells = summary_table.add_row().cells
            cells[0].text = str(label)
            cells[1].text = str(value)
            cells[0].paragraphs[0].runs[0].bold = True
        doc.add_paragraph()

    headers = list(report["headers"])
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = str(header)
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)

    for row in report["rows"]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = _text(value)
            cell.paragraphs[0].runs[0].font.size = Pt(9)

    if not report["rows"]:
        doc.add_paragraph("No records found.")

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

